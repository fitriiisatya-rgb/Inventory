#!/usr/bin/env python3
"""
PHASE G-DATA 2 — generates migration/templates/final_opening_stock_template.xlsx,
the format the owner/admin uses to supply FINAL VERIFIED opening stock for
Gudang Besar/SCM, Cibadak, and Karang Tengah. This is a template only —
no real data is embedded, and generating it does not import anything.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT_PATH = "/home/user/Inventory/migration/templates/final_opening_stock_template.xlsx"

COLUMNS = [
    "Cutoff Date", "Warehouse Code", "SKU", "Item Name", "Global Base Unit",
    "Opening Qty Base", "Unit Cost Base", "Opening Value", "Expiry Date",
    "Batch Reference", "Source", "Verification Status", "Approved By", "Notes",
]

EXAMPLE_ROWS = [
    ["2026-09-30", "SCM", "700107", "CONTOH ITEM (hapus baris ini)", "KG",
     125.5, 38728.67, "=F2*G2", "", "OPENING-SCM-001", "Final verified stock count",
     "Verified by stock count", "Nama Admin", "Contoh baris — hapus sebelum isi data asli"],
    ["2026-09-30", "CIBADAK", "700107", "CONTOH ITEM (hapus baris ini)", "KG",
     20.0, 38728.67, "=F3*G3", "", "OPENING-CBD-001", "Final verified stock count",
     "Verified by stock count", "Nama Admin", "Contoh baris — hapus sebelum isi data asli"],
]

wb = openpyxl.Workbook()

# ---------------------------------------------------------------------
# INSTRUCTIONS sheet
# ---------------------------------------------------------------------
ws_instr = wb.active
ws_instr.title = "INSTRUCTIONS"
instructions = [
    "FINAL OPENING STOCK TEMPLATE — Phase G-DATA 2",
    "",
    "This is the AUTHORITATIVE final opening stock source. It replaces the",
    "September opening/In-Out reconstruction, legacy FIFO batches, and the",
    "old dbinventory stock calculation entirely — none of those are used to",
    "post production opening. Only THIS file (once verified and approved)",
    "becomes the real opening balance.",
    "",
    "WAREHOUSES: use warehouse codes SCM (Gudang Besar), CIBADAK, KARANG_TENGAH.",
    "",
    "GLOBAL BASE UNIT POLICY:",
    "- Always enter Opening Qty Base and Unit Cost Base in the item's Global",
    "  Base Unit (e.g. KG, PCS, LTR) — the SAME unit the system already uses",
    "  for that SKU. You do NOT need an approved CARTON/PACK/PAIL conversion",
    "  to fill in this template; base-unit qty + cost is always sufficient.",
    "- If a SKU's purchase-unit conversion is not yet approved, opening still",
    "  works normally in the base unit. It is only TRANSACTING later in an",
    "  unapproved unit (e.g. posting a receipt in CARTON) that is blocked",
    "  (error UNIT_CONVERSION_NOT_APPROVED) until that conversion is approved.",
    "",
    "UNIT COST NORMALIZATION: Unit Cost Base must be the cost per ONE unit",
    "of the Global Base Unit — never a package/purchase-unit price. Example:",
    "  1 PCS = 15 KG, package price Rp580,930 per PCS",
    "  -> Unit Cost Base = 580930 / 15 = Rp38,728.67 per KG",
    "  Rp305,000/KG, if Global Base Unit = GR -> Unit Cost Base = 305 (per GR)",
    "",
    "OPENING VALUE column is DISPLAY / REFERENCE ONLY (shown here as a",
    "formula, Qty x Cost) — the server always recalculates it from Opening",
    "Qty Base x Unit Cost Base and never trusts an uploaded Opening Value.",
    "",
    "NEGATIVE QUANTITY IS ALWAYS REJECTED. No exceptions.",
    "A positive quantity requires a positive Unit Cost Base.",
    "",
    "CUTOFF DATE: all rows, across all three warehouses, should share ONE",
    "cutoff date. If your warehouses were counted on different dates, the",
    "import will be rejected and ask you to reconcile to a single date first.",
    "",
    "Item Name and Global Base Unit are REFERENCE columns only, cross-checked",
    "against the item master (a mismatch is rejected) but never authoritative",
    "by themselves — the SKU is what identifies the item.",
    "",
    "Source / Verification Status / Approved By / Notes are free-text audit",
    "columns describing where this figure came from and who verified it.",
    "",
    "Remove the example rows on the Template sheet before entering real data.",
]
for i, line in enumerate(instructions, start=1):
    cell = ws_instr.cell(row=i, column=1, value=line)
    if i == 1:
        cell.font = Font(bold=True, size=13)
ws_instr.column_dimensions["A"].width = 100

# ---------------------------------------------------------------------
# Template sheet
# ---------------------------------------------------------------------
ws = wb.create_sheet("Template")
header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
for col_idx, col_name in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=col_idx, value=col_name)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill

for row_idx, row in enumerate(EXAMPLE_ROWS, start=2):
    for col_idx, value in enumerate(row, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)

for col_idx, col_name in enumerate(COLUMNS, start=1):
    width = max(16, min(40, len(col_name) + 6))
    ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

dv_warehouse = DataValidation(type="list", formula1='"SCM,CIBADAK,KARANG_TENGAH"', allow_blank=False)
ws.add_data_validation(dv_warehouse)
dv_warehouse.add("B2:B1048576")

ws.freeze_panes = "A2"

wb.save(OUT_PATH)
print(f"Wrote {OUT_PATH}")
